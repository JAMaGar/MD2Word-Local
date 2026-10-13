# -*- coding: utf-8 -*-
"""
conversor.py - Conversao local de Markdown (.md) para Word (.docx / .doc / .pdf).

Tudo roda na maquina do usuario. Nenhum dado e enviado para a internet:
imagens remotas (http/https) NAO sao baixadas; aparecem como texto no documento.

Recursos suportados
-------------------
- Titulos (#..######) com estilos nativos do Word (Titulo 1..6, entram no Sumario)
- Negrito, italico, tachado, `codigo`, links clicaveis, <sub>, <sup>, <u>, <mark>, <br>
- Listas com marcadores, numeradas (respeita o numero inicial), aninhadas e de tarefas [ ] [x]
- Citacoes (>) aninhadas, blocos de codigo, linhas horizontais (---)
- Tabelas GFM com alinhamento por coluna e cabecalho repetido em cada pagina
- Notas de rodape [^1] como notas de rodape REAIS do Word
- Equacoes $...$ e $$...$$ como equacoes editaveis do Word (quando o Office esta instalado)
- Imagens locais (caminho relativo ao .md) e imagens embutidas em base64
- Front matter YAML (---) e ignorado
- Saida .doc e .pdf via Microsoft Word (ou LibreOffice, se instalado)
"""
from __future__ import annotations

import base64
import datetime as _dt
import getpass
import glob
import html as _html
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from markdown_it import MarkdownIt
from markdown_it.tree import SyntaxTreeNode
from mdit_py_plugins.dollarmath import dollarmath_plugin
from mdit_py_plugins.footnote import footnote_plugin
from mdit_py_plugins.front_matter import front_matter_plugin
from mdit_py_plugins.tasklists import tasklists_plugin

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Pt, RGBColor
from docx.text.paragraph import Paragraph

__version__ = "1.0.0"

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
FORMATOS_WORD = {"doc": 0, "pdf": 17}  # wdFormatDocument (97-2003) / wdFormatPDF


# --------------------------------------------------------------------------------------
# Opcoes
# --------------------------------------------------------------------------------------
@dataclass
class Opcoes:
    papel: str = "A4"                 # "A4" ou "Carta"
    margem_cm: float = 2.5
    fonte: str = "Calibri"
    tamanho_pt: float = 11.0
    fonte_codigo: str = "Consolas"
    modelo_docx: str = ""             # .docx opcional cujos estilos serao herdados
    equacoes_nativas: bool = True     # $...$ -> equacao do Word (requer MML2OMML.XSL do Office)
    idioma: str = "pt-BR"
    justificar: bool = False


# --------------------------------------------------------------------------------------
# Utilidades de XML
# --------------------------------------------------------------------------------------
# ordem dos filhos de w:pPr segundo o schema (CT_PPrBase) - o Word e exigente com a ordem
_PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku",
    "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
    "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
    "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]


def _insert_pPr_child(pPr, child):
    name = etree.QName(child).localname
    idx = _PPR_ORDER.index(name)
    for existing in list(pPr):
        ename = etree.QName(existing).localname
        if ename in _PPR_ORDER and _PPR_ORDER.index(ename) > idx:
            existing.addprevious(child)
            return
    pPr.append(child)


_RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
    "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
    "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
    "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
    "specVanish", "oMath",
]


def _insert_rPr_child(rPr, child):
    idx = _RPR_ORDER.index(etree.QName(child).localname)
    for existing in list(rPr):
        ename = etree.QName(existing).localname
        if ename in _RPR_ORDER and _RPR_ORDER.index(ename) > idx:
            existing.addprevious(child)
            return
    rPr.append(child)


def _shd(fill_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    return shd


def _p_border(paragraph, side: str, sz=6, space=4, color="BFBFBF"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        _insert_pPr_child(pPr, pBdr)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)
    # ordem: top, left, bottom, right, between, bar
    order = ["top", "left", "bottom", "right", "between", "bar"]
    for ex in list(pBdr):
        if order.index(etree.QName(ex).localname) > order.index(side):
            ex.addprevious(b)
            return
    pBdr.append(b)


def _p_shading(paragraph, fill_hex):
    _insert_pPr_child(paragraph._p.get_or_add_pPr(), _shd(fill_hex))


def _set_cell_shading(cell, fill_hex):
    cell._tc.get_or_add_tcPr().append(_shd(fill_hex))


def _set_table_borders(table, color="A6A6A6", sz=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        borders.append(b)
    before = {"tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd"}
    for child in list(tblPr):
        if etree.QName(child).localname not in before:
            child.addprevious(borders)
            return
    tblPr.append(borders)


def _repeat_header(row):
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(h)


# --------------------------------------------------------------------------------------
# Equacoes: LaTeX -> MathML -> OMML (usando o MML2OMML.XSL que acompanha o Microsoft Office)
# --------------------------------------------------------------------------------------
_XSLT_CACHE = {"loaded": False, "xslt": None, "path": None}


def localizar_mml2omml() -> str | None:
    candidatos = []
    for d in (Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)),
              Path(sys.executable).resolve().parent, Path(__file__).resolve().parent):
        candidatos += [d / "MML2OMML.XSL", d / "assets" / "MML2OMML.XSL"]
    for c in candidatos:
        if c.is_file():
            return str(c)
    if os.name == "nt":
        roots = {os.environ.get("ProgramFiles", r"C:\Program Files"),
                 os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
                 os.environ.get("ProgramW6432", r"C:\Program Files")}
        padroes = [r"Microsoft Office\root\Office*\MML2OMML.XSL",
                   r"Microsoft Office\Office*\MML2OMML.XSL"]
        for r in roots:
            if not r:
                continue
            for p in padroes:
                achados = sorted(glob.glob(os.path.join(r, p)), reverse=True)
                if achados:
                    return achados[0]
    return None


def _get_xslt():
    if not _XSLT_CACHE["loaded"]:
        _XSLT_CACHE["loaded"] = True
        path = localizar_mml2omml()
        if path:
            try:
                _XSLT_CACHE["xslt"] = etree.XSLT(etree.parse(path))
                _XSLT_CACHE["path"] = path
            except Exception:
                _XSLT_CACHE["xslt"] = None
    return _XSLT_CACHE["xslt"]


def latex_para_omml(latex: str, bloco: bool):
    """Retorna um elemento m:oMath (ou None se nao for possivel converter)."""
    xslt = _get_xslt()
    if xslt is None:
        return None
    try:
        from latex2mathml.converter import convert as l2m
        mathml = l2m(latex, display="block" if bloco else "inline")
        root = xslt(etree.fromstring(mathml.encode("utf-8"))).getroot()
        if root is None:
            return None
        if etree.QName(root).localname == "oMath":
            return root
        return root.find(f".//{{{M_NS}}}oMath")
    except Exception:
        return None


_LATEX_UNICODE = {
    r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\delta": "δ", r"\epsilon": "ε",
    r"\varepsilon": "ε", r"\zeta": "ζ", r"\eta": "η", r"\theta": "θ", r"\vartheta": "ϑ",
    r"\iota": "ι", r"\kappa": "κ", r"\lambda": "λ", r"\mu": "μ", r"\nu": "ν", r"\xi": "ξ",
    r"\pi": "π", r"\rho": "ρ", r"\sigma": "σ", r"\tau": "τ", r"\upsilon": "υ", r"\phi": "φ",
    r"\varphi": "φ", r"\chi": "χ", r"\psi": "ψ", r"\omega": "ω", r"\Gamma": "Γ",
    r"\Delta": "Δ", r"\Theta": "Θ", r"\Lambda": "Λ", r"\Xi": "Ξ", r"\Pi": "Π",
    r"\Sigma": "Σ", r"\Phi": "Φ", r"\Psi": "Ψ", r"\Omega": "Ω", r"\times": "×",
    r"\cdot": "·", r"\pm": "±", r"\mp": "∓", r"\leq": "≤", r"\le": "≤", r"\geq": "≥",
    r"\ge": "≥", r"\neq": "≠", r"\ne": "≠", r"\approx": "≈", r"\sim": "∼", r"\propto": "∝",
    r"\infty": "∞", r"\partial": "∂", r"\nabla": "∇", r"\sum": "∑", r"\prod": "∏",
    r"\int": "∫", r"\sqrt": "√", r"\rightarrow": "→", r"\to": "→", r"\leftarrow": "←",
    r"\Rightarrow": "⇒", r"\Leftrightarrow": "⇔", r"\leftrightarrow": "↔", r"\degree": "°",
    r"\circ": "∘", r"\in": "∈", r"\notin": "∉", r"\subset": "⊂", r"\cup": "∪", r"\cap": "∩",
    r"\forall": "∀", r"\exists": "∃", r"\hbar": "ħ", r"\ell": "ℓ", r"\AA": "Å",
    r"\langle": "⟨", r"\rangle": "⟩", r"\ldots": "…", r"\cdots": "⋯", r"\,": " ",
    r"\;": " ", r"\quad": "  ", r"\qquad": "    ", r"\%": "%", r"\left": "", r"\right": "",
}
_SUP_CH = "0123456789+-=()nix"
_SUB_CH = "0123456789+-=()aeoxijnmr"
_SUP = str.maketrans(_SUP_CH, "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱˣ")
_SUB = str.maketrans(_SUB_CH, "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑₒₓᵢⱼₙₘᵣ")


def latex_para_texto(latex: str) -> str:
    """Aproximacao em Unicode, usada quando nao ha conversao nativa de equacoes."""
    s = latex
    s = re.sub(r"\\(?:mathrm|text|mathbf|mathit|operatorname|textrm|mathsf)\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    for k in sorted(_LATEX_UNICODE, key=len, reverse=True):
        s = re.sub(re.escape(k) + r"(?![A-Za-z])", lambda _m, v=_LATEX_UNICODE[k]: v, s)

    def sup(m):
        t = m.group(1) if m.group(1) is not None else m.group(2)
        return t.translate(_SUP) if all(ch in _SUP_CH for ch in t) else "^(" + t + ")"

    def sub(m):
        t = m.group(1) if m.group(1) is not None else m.group(2)
        return t.translate(_SUB) if all(ch in _SUB_CH for ch in t) else "_(" + t + ")"

    s = re.sub(r"\^\{([^{}]*)\}|\^(\S)", sup, s)
    s = re.sub(r"_\{([^{}]*)\}|_(\S)", sub, s)
    return s.replace("{", "").replace("}", "").strip()


# --------------------------------------------------------------------------------------
# Numeracao de listas (numbering.xml)
# --------------------------------------------------------------------------------------
class Numeracao:
    BULLETS = ["•", "◦", "▪"] * 3
    FMT_NUM = ["decimal", "lowerLetter", "lowerRoman"] * 3

    def __init__(self, document):
        self.doc = document
        try:
            self.part = document.part.numbering_part
        except (NotImplementedError, KeyError):
            self.part = self._criar_parte()
        self.root = self.part.element
        ids = [int(x) for x in self.root.xpath("./w:abstractNum/@w:abstractNumId")]
        nids = [int(x) for x in self.root.xpath("./w:num/@w:numId")]
        self._next_abs = (max(ids) + 1) if ids else 1
        self._next_num = (max(nids) + 1) if nids else 1
        self.abs_bullet = self._abstract(ordenada=False)
        self.abs_num = self._abstract(ordenada=True)
        self.num_bullet = self._num(self.abs_bullet)

    def _criar_parte(self):
        from docx.parts.numbering import NumberingPart
        part = NumberingPart(
            PackURI("/word/numbering.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml",
            parse_xml("<w:numbering %s/>" % nsdecls("w")), self.doc.part.package)
        self.doc.part.relate_to(part, RT.NUMBERING)
        return part

    @staticmethod
    def indent_twips(ilvl: int) -> int:
        return 720 * (ilvl + 1)

    def _abstract(self, ordenada: bool) -> int:
        aid = self._next_abs
        self._next_abs += 1
        a = OxmlElement("w:abstractNum")
        a.set(qn("w:abstractNumId"), str(aid))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "hybridMultilevel")
        a.append(mlt)
        for lvl in range(9):
            L = OxmlElement("w:lvl")
            L.set(qn("w:ilvl"), str(lvl))
            st = OxmlElement("w:start"); st.set(qn("w:val"), "1"); L.append(st)
            nf = OxmlElement("w:numFmt")
            nf.set(qn("w:val"), self.FMT_NUM[lvl] if ordenada else "bullet")
            L.append(nf)
            lt = OxmlElement("w:lvlText")
            lt.set(qn("w:val"), f"%{lvl + 1}." if ordenada else self.BULLETS[lvl])
            L.append(lt)
            jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); L.append(jc)
            pPr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(self.indent_twips(lvl)))
            ind.set(qn("w:hanging"), "360")
            pPr.append(ind)
            L.append(pPr)
            if not ordenada:
                rPr = OxmlElement("w:rPr")
                f = OxmlElement("w:rFonts")
                for k in ("ascii", "hAnsi", "cs"):
                    f.set(qn(f"w:{k}"), "Arial")
                f.set(qn("w:hint"), "default")
                rPr.append(f)
                L.append(rPr)
            a.append(L)
        first_num = self.root.find(qn("w:num"))  # abstractNum precisa vir antes dos w:num
        if first_num is not None:
            first_num.addprevious(a)
        else:
            self.root.append(a)
        return aid

    def _num(self, abstract_id: int, start: int | None = None, ilvl: int = 0) -> int:
        nid = self._next_num
        self._next_num += 1
        n = OxmlElement("w:num")
        n.set(qn("w:numId"), str(nid))
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), str(abstract_id))
        n.append(an)
        if start is not None:
            ov = OxmlElement("w:lvlOverride")
            ov.set(qn("w:ilvl"), str(ilvl))
            so = OxmlElement("w:startOverride")
            so.set(qn("w:val"), str(start))
            ov.append(so)
            n.append(ov)
        self.root.append(n)
        return nid

    def nova_lista(self, ordenada: bool, start: int, ilvl: int) -> int:
        if not ordenada:
            return self.num_bullet
        return self._num(self.abs_num, start=start, ilvl=ilvl)

    @staticmethod
    def aplicar(paragraph, num_id: int, ilvl: int):
        numPr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), str(ilvl))
        ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num_id))
        numPr.append(il); numPr.append(ni)
        _insert_pPr_child(paragraph._p.get_or_add_pPr(), numPr)


# --------------------------------------------------------------------------------------
# Notas de rodape reais do Word (footnotes.xml)
# --------------------------------------------------------------------------------------
_FOOTNOTES_XML = (
    '<w:footnotes %s>'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
    'w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" '
    'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    '</w:footnotes>'
) % nsdecls("w", "r")

# filhos de w:settings que precedem w:footnotePr (CT_Settings)
_SETTINGS_BEFORE_FOOTNOTEPR = set("""
writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries
displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData
embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins
alignBordersAndEdges bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop
hideSpellingErrors hideGrammaticalErrors activeWritingStyle proofState formsDesign
attachedTemplate linkStyles stylePaneFormatFilter stylePaneSortMethod documentType mailMerge
revisionView trackRevisions doNotTrackMoves doNotTrackFormatting documentProtection
autoFormatOverride styleLockTheme styleLockQFSet defaultTabStop autoHyphenation
consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps showEnvelope summaryLength
clickAndTypeStyle defaultTableStyle evenAndOddHeaders bookFoldRevPrinting bookFoldPrinting
bookFoldPrintingSheets drawingGridHorizontalSpacing drawingGridVerticalSpacing
displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery
doNotUseMarginsForDrawingGridOrigin drawingGridHorizontalOrigin drawingGridVerticalOrigin
doNotShadeFormData noPunctuationKerning characterSpacingControl printTwoOnOne
strictFirstAndLastChars noLineBreaksAfter noLineBreaksBefore savePreviewPicture
doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent alwaysShowPlaceholderText
doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags
alwaysMergeEmptyNamespace updateFields hdrShapeDefaults
""".split())


class _Container:
    """Objeto minimo para que python-docx crie paragrafos dentro de uma nota de rodape."""

    def __init__(self, element, part):
        self._element = element
        self.part = part

    def add_paragraph(self):
        p = OxmlElement("w:p")
        self._element.append(p)
        return Paragraph(p, self)


class NotasRodape:
    def __init__(self, document):
        self.doc = document
        self.part = None
        self.root = None

    def _garantir(self):
        if self.part is not None:
            return
        self.root = parse_xml(_FOOTNOTES_XML)
        self.part = XmlPart(
            PackURI("/word/footnotes.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            self.root, self.doc.part.package)
        self.doc.part.relate_to(self.part, RT.FOOTNOTES)
        settings = self.doc.settings.element
        if settings.find(qn("w:footnotePr")) is None:
            fpr = parse_xml('<w:footnotePr %s><w:footnote w:id="-1"/><w:footnote w:id="0"/>'
                            '</w:footnotePr>' % nsdecls("w"))
            for child in settings:
                if etree.QName(child).localname not in _SETTINGS_BEFORE_FOOTNOTEPR:
                    child.addprevious(fpr)
                    break
            else:
                settings.append(fpr)
        styles = self.doc.styles
        nomes = {s.name.lower() for s in styles}
        if "footnote text" not in nomes:
            st = styles.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            st.font.size = Pt(9)
            st.paragraph_format.space_after = Pt(0)
            st.paragraph_format.line_spacing = 1.0
        if "footnote reference" not in nomes:
            st = styles.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
            st.font.superscript = True

    def nova(self, word_id: int) -> _Container:
        self._garantir()
        fn = OxmlElement("w:footnote")
        fn.set(qn("w:id"), str(word_id))
        self.root.append(fn)
        return _Container(fn, self.part)


def _style_id(doc, name, default):
    try:
        return doc.styles[name].style_id
    except KeyError:
        return default


# --------------------------------------------------------------------------------------
# Renderizador
# --------------------------------------------------------------------------------------
@dataclass
class _Ctx:
    container: object
    ilvl: int = -1            # nivel de lista atual (-1 = fora de lista)
    quote: int = 0            # profundidade de citacao
    num_id: int | None = None
    first_in_item: bool = False
    tight: bool = False
    in_footnote: bool = False


@dataclass
class _Fmt:
    bold: bool = False
    italic: bool = False
    strike: bool = False
    code: bool = False
    sub: bool = False
    sup: bool = False
    underline: bool = False
    mark: bool = False
    link: str | None = None


def _com(fmt: _Fmt, **kw) -> _Fmt:
    return _Fmt(**{**fmt.__dict__, **kw})


class Renderizador:
    def __init__(self, md_path: Path | None, opcoes: Opcoes, avisos: list):
        self.md_path = md_path
        self.base_dir = md_path.parent if md_path else Path.cwd()
        self.op = opcoes
        self.avisos = avisos
        self.doc = self._novo_documento()
        self.num = Numeracao(self.doc)
        self.notas = NotasRodape(self.doc)
        self.footnote_map: dict[int, int] = {}
        self.titulo = None
        self.html_fmt = _Fmt()
        self._hyperlink_el = None
        self._avisou_eq = False

    # ---------------------------------------------------------------- documento base
    def _novo_documento(self):
        op = self.op
        if op.modelo_docx and Path(op.modelo_docx).is_file():
            doc = Document(op.modelo_docx)
            body = doc.element.body
            for child in list(body):
                if etree.QName(child).localname != "sectPr":
                    body.remove(child)
            return doc
        doc = Document()
        sec = doc.sections[0]
        if op.papel.lower().startswith("a4"):
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        else:
            sec.page_width, sec.page_height = Emu(7772400), Emu(10058400)
        sec.orientation = WD_ORIENT.PORTRAIT
        m = Cm(op.margem_cm)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = m
        normal = doc.styles["Normal"]
        normal.font.size = Pt(op.tamanho_pt)
        self._set_style_font(normal, op.fonte)
        rpr = normal.element.get_or_add_rPr()
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:val"), op.idioma)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for i in range(1, 7):
            try:
                hs = doc.styles[f"Heading {i}"]
            except KeyError:
                continue
            self._set_style_font(hs, op.fonte)
            hs.paragraph_format.space_before = Pt([18, 14, 12, 10, 10, 10][i - 1])
            hs.paragraph_format.space_after = Pt(6)
        return doc

    @staticmethod
    def _set_style_font(style, fonte):
        style.font.name = fonte
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rf.attrib.pop(qn(f"w:{attr}"), None)
        for k in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn(f"w:{k}"), fonte)

    def _usable_width(self):
        sec = self.doc.sections[-1]
        return sec.page_width - sec.left_margin - sec.right_margin

    # ---------------------------------------------------------------- paragrafos
    def _new_par(self, ctx: _Ctx, style=None):
        p = ctx.container.add_paragraph()
        if ctx.in_footnote:
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), _style_id(self.doc, "Footnote Text", "FootnoteText"))
            _insert_pPr_child(p._p.get_or_add_pPr(), ps)
        elif style:
            try:
                p.style = self.doc.styles[style]
            except KeyError:
                pass
        fmt = p.paragraph_format
        left = 0
        if ctx.ilvl >= 0:
            if ctx.first_in_item and ctx.num_id is not None:
                Numeracao.aplicar(p, ctx.num_id, ctx.ilvl)
                ctx.first_in_item = False
            else:
                left = Numeracao.indent_twips(ctx.ilvl)
        if ctx.quote:
            left += 425 * ctx.quote
            _p_border(p, "left", sz=18, space=8, color="A6A6A6")
        if left:
            fmt.left_indent = Emu(left * 635)  # 1 twip = 635 EMU
        if ctx.ilvl >= 0 and ctx.tight:
            fmt.space_after = Pt(2)
        if self.op.justificar and not ctx.in_footnote and style is None and ctx.ilvl < 0:
            fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    # ---------------------------------------------------------------- entrada principal
    def render(self, texto_md: str):
        md = (
            MarkdownIt("commonmark", {"html": True, "typographer": False})
            .enable(["table", "strikethrough"])
            .use(front_matter_plugin)
            .use(tasklists_plugin)
            .use(footnote_plugin)
            .use(dollarmath_plugin, allow_space=False, allow_digits=False, double_inline=True)
        )
        tree = SyntaxTreeNode(md.parse(texto_md))
        ctx = _Ctx(container=self.doc)

        # 1) notas de rodape primeiro (as referencias aparecem antes, no corpo)
        for node in tree.children:
            if node.type != "footnote_block":
                continue
            for fn in node.children:
                if fn.type != "footnote":
                    continue
                word_id = len(self.footnote_map) + 1
                self.footnote_map[fn.meta.get("id")] = word_id
                fctx = _Ctx(container=self.notas.nova(word_id), in_footnote=True)
                first = True
                for child in fn.children:
                    if child.type == "paragraph" and first:
                        p = self._new_par(fctx)
                        self._footnote_mark(p, ref=False)
                        p.add_run(" ")
                        self._inline(p, child.children[0] if child.children else None, fctx)
                        first = False
                    elif child.type != "footnote_anchor":
                        self._block(child, fctx)
                if first:
                    self._footnote_mark(self._new_par(fctx), ref=False)

        # 2) corpo
        for node in tree.children:
            if node.type != "footnote_block":
                self._block(node, ctx)

        cp = self.doc.core_properties
        cp.title = self.titulo or (self.md_path.stem if self.md_path else "Documento")
        try:
            cp.author = getpass.getuser()
        except Exception:
            cp.author = ""
        cp.last_modified_by = cp.author
        cp.comments = cp.subject = cp.keywords = ""
        now = _dt.datetime.now()
        cp.created = cp.modified = now
        cp.revision = 1
        return self.doc

    # ---------------------------------------------------------------- blocos
    def _block(self, node, ctx: _Ctx):
        t = node.type
        inline = node.children[0] if node.children and node.children[0].type == "inline" else None
        if t == "heading":
            level = int(node.tag[1])
            if ctx.in_footnote or ctx.ilvl >= 0 or ctx.quote:
                self._inline(self._new_par(ctx), inline, ctx, _Fmt(bold=True))
                return
            p = self._new_par(ctx, style=f"Heading {level}")
            self._inline(p, inline, ctx)
            if level == 1 and self.titulo is None:
                self.titulo = p.text
        elif t == "paragraph":
            p = self._new_par(ctx)
            self._inline(p, inline, ctx)
            if inline is not None and len(inline.children) == 1 and inline.children[0].type == "image":
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif t in ("bullet_list", "ordered_list"):
            self._list(node, ctx)
        elif t == "blockquote":
            sub = _Ctx(**{**ctx.__dict__, "quote": ctx.quote + 1})
            for child in node.children:
                self._block(child, sub)
            ctx.first_in_item = sub.first_in_item
        elif t in ("fence", "code_block"):
            self._code(node, ctx)
        elif t == "hr":
            p = self._new_par(ctx)
            _p_border(p, "bottom", sz=6, space=1, color="A6A6A6")
            p.paragraph_format.space_after = Pt(12)
        elif t == "table":
            self._table(node, ctx)
        elif t in ("math_block", "math_block_label"):
            self._math_block(node.content, ctx)
        elif t == "html_block":
            self._html_block(node.content, ctx)
        elif t in ("front_matter", "footnote_block", "footnote_anchor"):
            pass
        else:
            for child in node.children or []:
                self._block(child, ctx)

    def _list(self, node, ctx: _Ctx):
        ordenada = node.type == "ordered_list"
        ilvl = min(ctx.ilvl + 1, 8)
        start = 1
        if ordenada:
            try:
                start = int(node.attrs.get("start", 1))
            except (TypeError, ValueError):
                start = 1
        num_id = self.num.nova_lista(ordenada, start, ilvl)
        # markdown-it marca os paragrafos de listas compactas como "hidden"
        tight = all(getattr(g, "hidden", False)
                    for item in node.children for g in (item.children or [])
                    if g.type == "paragraph")
        for item in node.children:
            sub = _Ctx(container=ctx.container, ilvl=ilvl, quote=ctx.quote, num_id=num_id,
                       first_in_item=True, tight=tight, in_footnote=ctx.in_footnote)
            if not item.children:
                self._new_par(sub)
                continue
            for child in item.children:
                self._block(child, sub)

    def _code(self, node, ctx: _Ctx):
        lines = node.content.rstrip("\n").split("\n")
        p = self._new_par(ctx)
        for side in ("top", "left", "bottom", "right"):
            _p_border(p, side, sz=4, space=4, color="D9D9D9")
        _p_shading(p, "F5F5F5")
        pf = p.paragraph_format
        pf.space_before = Pt(4)
        pf.space_after = Pt(8)
        pf.line_spacing = 1.0
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        size = max(self.op.tamanho_pt - 1.5, 7)
        for i, line in enumerate(lines):
            for j, part in enumerate(line.split("\t")):
                if j:
                    p.add_run().add_tab()
                self._fmt_code_run(p.add_run(part), size=size)
            if i < len(lines) - 1:
                p.add_run().add_break(WD_BREAK.LINE)

    def _fmt_code_run(self, r, size=None, shade=False):
        rPr = r._r.get_or_add_rPr()
        r.font.name = self.op.fonte_codigo
        rf = rPr.find(qn("w:rFonts"))
        for k in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn(f"w:{k}"), self.op.fonte_codigo)
        if size:
            r.font.size = Pt(size)
        if shade:
            _insert_rPr_child(rPr, _shd("EDEDED"))
        r.font.no_proof = True

    def _table(self, node, ctx: _Ctx):
        rows = []
        for section in node.children:  # thead / tbody
            for tr in section.children:
                cells = [c for c in tr.children if c.type in ("th", "td")]
                rows.append((section.type == "thead", cells))
        if not rows:
            return
        ncols = max(len(c) for _, c in rows)
        if ctx.in_footnote or not hasattr(ctx.container, "add_table"):
            for _, cells in rows:
                p = self._new_par(ctx)
                for k, c in enumerate(cells):
                    if k:
                        p.add_run(" | ")
                    self._inline(p, c.children[0] if c.children else None, ctx)
            return
        if ctx.container is self.doc:
            table = self.doc.add_table(rows=len(rows), cols=ncols)
        else:
            table = ctx.container.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(table)
        for r_i, (is_head, cells) in enumerate(rows):
            row = table.rows[r_i]
            if is_head:
                _repeat_header(row)
            for c_i in range(ncols):
                cell = row.cells[c_i]
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                if c_i < len(cells):
                    c = cells[c_i]
                    m = re.search(r"text-align:\s*(left|right|center)", str(c.attrs.get("style") or ""))
                    if m:
                        par.paragraph_format.alignment = {
                            "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                            "center": WD_ALIGN_PARAGRAPH.CENTER}[m.group(1)]
                    self._inline(par, c.children[0] if c.children else None, _Ctx(container=cell),
                                 _Fmt(bold=True) if is_head else None)
                if is_head:
                    _set_cell_shading(cell, "DCE6F2")
        self._new_par(ctx).paragraph_format.space_after = Pt(2)

    def _math_block(self, latex: str, ctx: _Ctx):
        p = self._new_par(ctx)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        omml = latex_para_omml(latex.strip(), bloco=True) if self.op.equacoes_nativas else None
        if omml is not None:
            etree.SubElement(p._p, f"{{{M_NS}}}oMathPara").append(omml)
        else:
            self._aviso_eq()
            r = p.add_run(latex_para_texto(latex))
            r.font.name = "Cambria Math"
            r.italic = True

    def _aviso_eq(self):
        if self.op.equacoes_nativas and not self._avisou_eq:
            self._avisou_eq = True
            if _get_xslt() is None:
                self.avisos.append("Equacoes inseridas como texto: o arquivo MML2OMML.XSL do "
                                   "Microsoft Office nao foi encontrado.")
            else:
                self.avisos.append("Alguma equacao nao pode ser convertida e foi inserida como texto.")

    def _html_block(self, html: str, ctx: _Ctx):
        s = html.strip()
        if re.search(r"page-break|pagebreak|<!--\s*(newpage|pagebreak)\s*-->", s, flags=re.I):
            self._new_par(ctx).add_run().add_break(WD_BREAK.PAGE)
            return
        if not s or re.fullmatch(r"<!--.*?-->", s, flags=re.S):
            return
        text = re.sub(r"<br\s*/?>", "\n", s, flags=re.I)
        text = re.sub(r"</(p|div|li|h\d|tr)>", "\n", text, flags=re.I)
        text = _html.unescape(re.sub(r"<[^>]+>", "", text)).strip()
        for line in text.split("\n"):
            if line.strip():
                self._new_par(ctx).add_run(line.strip())

    # ---------------------------------------------------------------- inline
    def _inline(self, paragraph, inline_node, ctx: _Ctx, base: _Fmt | None = None):
        if inline_node is None:
            return
        self.html_fmt = _Fmt()
        self._hyperlink_el = None
        self._walk(paragraph, inline_node.children, base or _Fmt(), ctx)

    def _walk(self, p, nodes, fmt: _Fmt, ctx: _Ctx):
        for n in nodes:
            t = n.type
            if t == "text":
                self._run(p, n.content, fmt)
            elif t == "softbreak":
                self._run(p, " ", fmt)
            elif t == "hardbreak":
                self._run(p, "", fmt).add_break(WD_BREAK.LINE)
            elif t == "strong":
                self._walk(p, n.children, _com(fmt, bold=True), ctx)
            elif t == "em":
                self._walk(p, n.children, _com(fmt, italic=True), ctx)
            elif t == "s":
                self._walk(p, n.children, _com(fmt, strike=True), ctx)
            elif t == "code_inline":
                self._run(p, n.content, _com(fmt, code=True))
            elif t == "link":
                self._link(p, n.attrs.get("href", ""), n.children, fmt, ctx)
            elif t == "image":
                self._image(p, n, ctx)
            elif t in ("math_inline", "math_inline_double"):
                self._math_inline(p, n.content, fmt)
            elif t == "footnote_ref":
                self._footnote_ref(p, n, ctx)
            elif t == "html_inline":
                self._html_inline(p, n.content, fmt)
            elif n.children:
                self._walk(p, n.children, fmt, ctx)
            elif n.content:
                self._run(p, n.content, fmt)

    _HTML_MAP = {"sub": "sub", "sup": "sup", "u": "underline", "ins": "underline",
                 "b": "bold", "strong": "bold", "i": "italic", "em": "italic",
                 "mark": "mark", "s": "strike", "del": "strike", "strike": "strike",
                 "kbd": "code", "code": "code", "tt": "code"}

    def _html_inline(self, p, html: str, fmt: _Fmt):
        h = html.strip().lower()
        if re.fullmatch(r"<br\s*/?>", h):
            self._run(p, "", fmt).add_break(WD_BREAK.LINE)
            return
        if "checkbox" in h:
            self._run(p, "☑" if "checked" in h else "☐", fmt)
            return
        m = re.fullmatch(r"<(/?)([a-z0-9]+)[^>]*>", h)
        if m and m.group(2) in self._HTML_MAP:
            setattr(self.html_fmt, self._HTML_MAP[m.group(2)], m.group(1) != "/")

    def _run(self, p, text, fmt: _Fmt):
        r = p.add_run(text)
        h = self.html_fmt
        if fmt.bold or h.bold:
            r.bold = True
        if fmt.italic or h.italic:
            r.italic = True
        if fmt.strike or h.strike:
            r.font.strike = True
        if fmt.underline or h.underline:
            r.underline = True
        if fmt.sub or h.sub:
            r.font.subscript = True
        if fmt.sup or h.sup:
            r.font.superscript = True
        if fmt.mark or h.mark:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if fmt.code or h.code:
            self._fmt_code_run(r, shade=True)
            r.font.color.rgb = RGBColor(0xA3, 0x1F, 0x4A)
        if fmt.link:
            r.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            r.underline = True
            if self._hyperlink_el is not None:
                self._hyperlink_el.append(r._r)
        return r

    def _link(self, p, href, children, fmt: _Fmt, ctx: _Ctx):
        if not href:
            self._walk(p, children, fmt, ctx)
            return
        hl = OxmlElement("w:hyperlink")
        if href.startswith("#"):
            hl.set(qn("w:anchor"), href[1:])
        else:
            alvo_href = href
            if not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", href) or re.match(r"^[a-zA-Z]:[\\/]", href):
                alvo = Path(href) if Path(href).is_absolute() else (self.base_dir / href)
                if alvo.exists():
                    alvo_href = alvo.resolve().as_uri()
            try:
                hl.set(qn("r:id"), p.part.relate_to(alvo_href, RT.HYPERLINK, is_external=True))
            except Exception:
                self._walk(p, children, _com(fmt, underline=True), ctx)
                return
        p._p.append(hl)
        prev = self._hyperlink_el
        self._hyperlink_el = hl
        self._walk(p, children, _com(fmt, link=href), ctx)
        self._hyperlink_el = prev
        if len(hl) == 0:
            p._p.remove(hl)

    def _image(self, p, node, ctx: _Ctx):
        src = node.attrs.get("src", "") or ""
        alt = node.content or ""
        stream = None
        if src.startswith("data:image/"):
            try:
                stream = io.BytesIO(base64.b64decode(src.split(",", 1)[1]))
            except Exception:
                self.avisos.append("Imagem base64 invalida ignorada.")
        elif re.match(r"^https?://", src, flags=re.I):
            self.avisos.append(f"Imagem da internet nao foi baixada (modo offline): {src}")
        else:
            from urllib.parse import unquote
            path = Path(unquote(re.sub(r"^file:/+", "", src)))
            if not path.is_absolute():
                path = self.base_dir / path
            exts = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".emf", ".wmf")
            if path.is_file() and path.suffix.lower() in exts:
                stream = str(path)
            elif path.is_file():
                self.avisos.append(f"Formato de imagem nao suportado pelo Word: {path.name}")
            else:
                self.avisos.append(f"Imagem nao encontrada: {src}")
        if stream is None or ctx.in_footnote:
            r = self._run(p, f"[Imagem: {alt or src}]", _Fmt(italic=True))
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            return
        try:
            run = p.add_run()
            pic = run.add_picture(stream)
            maxw = self._usable_width()
            if ctx.ilvl >= 0:
                maxw -= Emu(Numeracao.indent_twips(ctx.ilvl) * 635)
            if pic.width > maxw:
                ratio = maxw / pic.width
                pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
            if alt:
                docPr = run._r.find(".//{http://schemas.openxmlformats.org/drawingml/2006/"
                                    "wordprocessingDrawing}docPr")
                if docPr is not None:
                    docPr.set("descr", alt)
        except Exception as e:
            self.avisos.append(f"Falha ao inserir imagem {src}: {e}")
            self._run(p, f"[Imagem: {alt or src}]", _Fmt(italic=True))

    def _math_inline(self, p, latex, fmt):
        omml = latex_para_omml(latex.strip(), bloco=False) if self.op.equacoes_nativas else None
        if omml is not None:
            p._p.append(omml)
        else:
            self._aviso_eq()
            self._run(p, latex_para_texto(latex), _com(fmt, italic=True)).font.name = "Cambria Math"

    def _footnote_ref(self, p, node, ctx):
        word_id = self.footnote_map.get(node.meta.get("id"))
        if word_id is None or ctx.in_footnote:
            r = p.add_run(f"[{(node.meta.get('id') or 0) + 1}]")
            r.font.superscript = True
            return
        self._footnote_mark(p, ref=True, word_id=word_id)

    def _footnote_mark(self, p, ref: bool, word_id: int = 0):
        r = p.add_run()
        rPr = r._r.get_or_add_rPr()
        rs = OxmlElement("w:rStyle")
        rs.set(qn("w:val"), _style_id(self.doc, "Footnote Reference", "FootnoteReference"))
        rPr.insert(0, rs)
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript")
        rPr.append(va)
        if ref:
            el = OxmlElement("w:footnoteReference")
            el.set(qn("w:id"), str(word_id))
        else:
            el = OxmlElement("w:footnoteRef")
        r._r.append(el)


# --------------------------------------------------------------------------------------
# API publica
# --------------------------------------------------------------------------------------
def ler_markdown(path) -> str:
    data = Path(path).read_bytes()
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def markdown_para_docx(texto_md: str, destino, opcoes: Opcoes | None = None, md_path=None) -> list:
    """Converte o texto Markdown e salva em `destino` (.docx). Retorna a lista de avisos."""
    avisos: list[str] = []
    r = Renderizador(Path(md_path) if md_path else None, opcoes or Opcoes(), avisos)
    r.render(texto_md).save(str(destino))
    return avisos


def _caminho_livre(p: Path) -> Path:
    i = 1
    c = p
    while c.exists():
        c = p.with_name(f"{p.stem} ({i}){p.suffix}")
        i += 1
    return c


def _gravavel(p: Path) -> bool:
    if not p.exists():
        return True
    try:
        with open(p, "ab"):
            return True
    except OSError:
        return False


def localizar_libreoffice() -> str | None:
    s = shutil.which("soffice") or shutil.which("libreoffice")
    if s:
        return s
    if os.name == "nt":
        for r in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                  os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            c = Path(r) / "LibreOffice" / "program" / "soffice.exe"
            if c.is_file():
                return str(c)
    return None


# Script PowerShell que usa o proprio Word (automacao COM) para salvar .doc/.pdf.
# Caminhos passam por variaveis de ambiente para evitar problemas com aspas/acentos.
_PS_WORD = r"""
$ErrorActionPreference = 'Stop'
$w = $null
try {
  $w = New-Object -ComObject Word.Application
  $w.Visible = $false
  $w.DisplayAlerts = 0
  $d = $w.Documents.Open([string]$env:MD2W_IN, $false, $true, $false)
  if ([int]$env:MD2W_FMT -eq 17) { $d.ExportAsFixedFormat([string]$env:MD2W_OUT, 17) }
  else { $d.SaveAs2([string]$env:MD2W_OUT, [int]$env:MD2W_FMT) }
  $d.Close(0)
  exit 0
} catch {
  [Console]::Error.WriteLine($_.Exception.Message)
  exit 3
} finally {
  if ($w -ne $null) { $w.Quit(); [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($w) }
}
"""


def word_disponivel() -> bool:
    if os.name != "nt":
        return False
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Word.Application\CLSID"))
        return True
    except Exception:
        return False


def converter_via_word(docx, destino, formato: str, timeout=180):
    env = dict(os.environ)
    env["MD2W_IN"] = str(Path(docx).resolve())
    env["MD2W_OUT"] = str(Path(destino).resolve())
    env["MD2W_FMT"] = str(FORMATOS_WORD[formato])
    res = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
         "-Command", _PS_WORD],
        env=env, capture_output=True, text=True, timeout=timeout,
        creationflags=0x08000000 if os.name == "nt" else 0)  # CREATE_NO_WINDOW
    if res.returncode != 0 or not Path(destino).exists():
        raise RuntimeError((res.stderr or res.stdout or "falha no Word").strip()[:300])


def converter_via_libreoffice(docx, destino, formato: str, timeout=180):
    soffice = localizar_libreoffice()
    if not soffice:
        raise RuntimeError("LibreOffice nao encontrado")
    with tempfile.TemporaryDirectory() as td:
        filtro = {"doc": "doc:MS Word 97", "pdf": "pdf:writer_pdf_Export"}[formato]
        subprocess.run([soffice, "--headless", "--norestore", "--convert-to", filtro,
                        "--outdir", td, str(docx)], capture_output=True, timeout=timeout,
                       creationflags=0x08000000 if os.name == "nt" else 0)
        out = Path(td) / (Path(docx).stem + "." + formato)
        if not out.exists():
            raise RuntimeError("LibreOffice nao gerou o arquivo")
        shutil.move(str(out), str(destino))


def converter_arquivo(md_path, pasta_saida=None, formatos=("docx",), opcoes: Opcoes | None = None,
                      sobrescrever=True, log=print) -> list:
    """Converte um arquivo .md. Retorna a lista de arquivos gerados (Path)."""
    md_path = Path(md_path)
    pasta = Path(pasta_saida) if pasta_saida else md_path.parent
    pasta.mkdir(parents=True, exist_ok=True)
    texto = ler_markdown(md_path)
    formatos = [f.lower() for f in formatos] or ["docx"]
    gerados = []

    def destino_para(ext):
        d = pasta / f"{md_path.stem}.{ext}"
        if not sobrescrever:
            return _caminho_livre(d)
        if not _gravavel(d):
            novo = _caminho_livre(d)
            log(f"  ! {d.name} esta aberto em outro programa; salvando como {novo.name}")
            return novo
        return d

    tmpdir = None
    if "docx" in formatos:
        docx_path = destino_para("docx")
    else:
        tmpdir = tempfile.mkdtemp(prefix="md2word_")
        docx_path = Path(tmpdir) / f"{md_path.stem}.docx"
    try:
        for a in dict.fromkeys(markdown_para_docx(texto, docx_path, opcoes, md_path=md_path)):
            log(f"  ! {a}")
        if "docx" in formatos:
            gerados.append(docx_path)
        for fmt in ("doc", "pdf"):
            if fmt not in formatos:
                continue
            dest = destino_para(fmt)
            erro = ""
            if word_disponivel():
                try:
                    converter_via_word(docx_path, dest, fmt)
                    gerados.append(dest)
                    continue
                except Exception as e:
                    erro = f"(Word: {e})"
            try:
                converter_via_libreoffice(docx_path, dest, fmt)
                gerados.append(dest)
            except Exception as e:
                log(f"  X Nao foi possivel gerar .{fmt}: requer Microsoft Word ou LibreOffice "
                    f"instalado. {erro} {e}".strip())
    finally:
        if tmpdir:
            shutil.rmtree(tmpdir, ignore_errors=True)
    return gerados


def converter_texto(texto_md: str, destino, opcoes: Opcoes | None = None, log=print) -> Path:
    """Converte texto Markdown (ex.: colado da area de transferencia) para .docx, .doc ou .pdf,
    conforme a extensao de `destino`. Imagens relativas sao procuradas na pasta de destino."""
    destino = Path(destino)
    ext = destino.suffix.lower().lstrip(".") or "docx"
    md_virtual = destino.with_suffix(".md")
    if ext == "docx":
        for a in dict.fromkeys(markdown_para_docx(texto_md, destino, opcoes, md_path=md_virtual)):
            log(f"  ! {a}")
        return destino
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / (destino.stem + ".docx")
        for a in dict.fromkeys(markdown_para_docx(texto_md, tmp, opcoes, md_path=md_virtual)):
            log(f"  ! {a}")
        if word_disponivel():
            try:
                converter_via_word(tmp, destino, ext)
                return destino
            except Exception as e:
                log(f"  ! Word: {e}")
        converter_via_libreoffice(tmp, destino, ext)
        return destino


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description="Converte Markdown (.md) em Word, localmente.")
    ap.add_argument("arquivos", nargs="+", help="arquivos .md")
    ap.add_argument("-o", "--saida", help="pasta de saida (padrao: a mesma do .md)")
    ap.add_argument("-f", "--formato", action="append", choices=["docx", "doc", "pdf"],
                    help="formato de saida; pode repetir (padrao: docx)")
    ap.add_argument("--modelo", help=".docx modelo cujos estilos serao usados")
    ap.add_argument("--carta", action="store_true", help="papel Carta (padrao: A4)")
    ap.add_argument("--fonte", default="Calibri")
    ap.add_argument("--tamanho", type=float, default=11.0)
    ap.add_argument("--justificar", action="store_true")
    ap.add_argument("--sem-equacoes-nativas", action="store_true")
    a = ap.parse_args(argv)
    op = Opcoes(papel="Carta" if a.carta else "A4", fonte=a.fonte, tamanho_pt=a.tamanho,
                modelo_docx=a.modelo or "", equacoes_nativas=not a.sem_equacoes_nativas,
                justificar=a.justificar)
    for f in a.arquivos:
        print(f"> {f}")
        for g in converter_arquivo(f, a.saida, a.formato or ["docx"], op):
            print(f"  OK {g}")


if __name__ == "__main__":
    main()
